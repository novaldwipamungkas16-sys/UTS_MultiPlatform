from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

root=Path(r'D:/KULIAH_UIN_MALANG/SMT 5/OPT MULTIPLATFORM/UTS')
app=root/'fintrack_app'
out=root/'Laporan_UTS_Pemrograman_Multiplatform_Noval_Dwi_Pamungkas_240605110040.docx'
doc=Document()
sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.left_margin=Cm(4); sec.right_margin=Cm(3); sec.top_margin=Cm(4); sec.bottom_margin=Cm(3)
styles=doc.styles
styles['Normal'].font.name='Times New Roman'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); styles['Normal'].font.size=Pt(12)
styles['Normal'].paragraph_format.line_spacing=1.5; styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; styles['Normal'].paragraph_format.first_line_indent=Cm(1.25)
for s in ['Heading 1','Heading 2']:
 styles[s].font.name='Times New Roman'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); styles[s].font.bold=True; styles[s].font.size=Pt(12)

def p(t='', align=None, bold=False):
 x=doc.add_paragraph(); x.paragraph_format.line_spacing=1.5
 if align is not None: x.alignment=align
 r=x.add_run(t); r.bold=bold; r.font.name='Times New Roman'; r.font.size=Pt(12); return x
def h(t,n=1): doc.add_heading(t,level=n)
def caption(t): p(t,WD_ALIGN_PARAGRAPH.CENTER)

for t in ['LAPORAN UJIAN TENGAH SEMESTER (UTS)','MATA KULIAH PEMROGRAMAN MULTIPLATFORM','PENGEMBANGAN APLIKASI FINTRACK','CATATAN KEUANGAN PRIBADI BERBASIS FLUTTER','DENGAN ASYNCHRONOUS PROGRAMMING, SQLITE, DAN SHARED PREFERENCES']:
 p(t,WD_ALIGN_PARAGRAPH.CENTER,True)
p('Disusun Oleh:',WD_ALIGN_PARAGRAPH.CENTER); p('NOVAL DWI PAMUNGKAS',WD_ALIGN_PARAGRAPH.CENTER,True); p('NIM: 240605110040',WD_ALIGN_PARAGRAPH.CENTER); p('Kelas A — Semester 5',WD_ALIGN_PARAGRAPH.CENTER); p('Dosen Pengampu: A\'la Syauqi, M.Kom',WD_ALIGN_PARAGRAPH.CENTER)
p('PROGRAM STUDI S1 TEKNIK INFORMATIKA\nFAKULTAS SAINS DAN TEKNOLOGI\nUNIVERSITAS ISLAM NEGERI MAULANA MALIK IBRAHIM MALANG\n2026',WD_ALIGN_PARAGRAPH.CENTER,True)
doc.add_page_break()
h('BAB I'); h('PENDAHULUAN',2)
p('FinTrack merupakan aplikasi pencatatan keuangan pribadi berbasis Flutter dan Dart. Aplikasi mendukung pencatatan pemasukan dan pengeluaran, saldo bersih, anggaran, statistik kategori, mode gelap, serta preferensi pengguna. Data transaksi memakai SQLite, sedangkan preferensi memakai SharedPreferences. Aplikasi dirancang offline-first.')
h('Tujuan',2)
for t in ['Membangun aplikasi multiplatform dengan Flutter.','Menerapkan Future, async, await, Future.wait, dan Future.delayed.','Menerapkan CRUD SQLite dan SharedPreferences.','Menyediakan antarmuka responsif untuk pengelolaan keuangan.']: p(t, bold=False).style='List Bullet'
h('BAB II'); h('PEMBAHASAN',2)
h('Struktur Project',2)
for t in ['lib/models/transaction_model.dart — model transaksi.','lib/services/database_helper.dart — SQLite, CRUD, agregasi.','lib/services/preference_service.dart — preferensi pengguna.','lib/services/quote_service.dart — simulasi asynchronous.','lib/screens/ — dashboard, form transaksi, statistik, pengaturan.','lib/widgets/ — SummaryCard dan TransactionCard.']: p(t).style='List Bullet'
h('Implementasi Asynchronous dan Persistence',2)
p('HomeScreen memuat preferensi, transaksi, dan tips finansial secara paralel menggunakan Future.wait. QuoteService memakai Future.delayed selama 800 milidetik. SQLite menyimpan tabel transactions dengan kolom id, title, amount, type, category, date, dan note. SharedPreferences menyimpan is_dark_mode, user_name, dan monthly_budget.')
h('Pengujian Aktual',2)
p('Pengujian dijalankan pada project fintrack_app menggunakan Flutter SDK dari D:/src/flutter. Hasil: flutter analyze selesai tanpa masalah; flutter test berhasil dengan 10 test; flutter build apk --debug berhasil menghasilkan build/app/outputs/flutter-apk/app-debug.apk. Perangkat Android tidak terhubung saat pengujian, sehingga pengujian UI langsung pada Android dan screenshot baru dari emulator belum dapat dilakukan. Lima screenshot pada bagian berikut merupakan aset screenshot project yang diperiksa dimensinya dan dipakai sebagai dokumentasi halaman aplikasi.')
table=doc.add_table(rows=1, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'; table.rows[0].cells[0].text='Pemeriksaan'; table.rows[0].cells[1].text='Hasil'
for a,b in [('flutter analyze','Berhasil, tidak ada masalah'),('flutter test','Berhasil, 10 test passed'),('flutter build apk --debug','Berhasil, APK debug dibuat'),('Android device/AVD','Tidak tersedia saat pengujian')]:
 c=table.add_row().cells; c[0].text=a; c[1].text=b
h('Dokumentasi Tampilan',2)
imgs=[('01_dashboard_light.png','Gambar 2.1. Dashboard mode terang.'),('02_dashboard_dark.png','Gambar 2.2. Dashboard mode gelap.'),('03_form_tambah_transaksi.png','Gambar 2.3. Form tambah transaksi.'),('04_statistik_kategori.png','Gambar 2.4. Statistik kategori.'),('05_pengaturan_preferensi.png','Gambar 2.5. Pengaturan dan preferensi.')]
for fn,cap in imgs:
 path=app/'screenshots'/fn
 if path.exists(): doc.add_picture(str(path),width=Cm(11)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; caption(cap)
h('BAB III'); h('PENUTUP',2)
p('FinTrack berhasil dianalisis, diuji melalui test suite, dan dibangun menjadi APK debug. Fitur utama mencakup CRUD transaksi, persistence SQLite, preferensi SharedPreferences, statistik kategori, tema gelap, dan proses asynchronous. Validasi perangkat Android langsung masih menjadi langkah lanjutan karena tidak ada emulator atau perangkat Android aktif pada sesi pengujian.')
doc.save(out)
print(out)
